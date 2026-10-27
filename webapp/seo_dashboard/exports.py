"""Excel exports with native charts (openpyxl).

* ``full_report_xlsx``: every section of an audit on its own sheet, with a
  chart summarizing each part (scores, severities, status codes, keyword
  intents, ranking trends, CWV, AIO scores, events, roadmap).
* ``internal_links_xlsx``: the internal link building workbook (suggestions
  with source URL, target URL, anchor and context; orphans; inlink
  distribution; links to redirects; broken internal links).
"""

from __future__ import annotations

import io
from collections import Counter

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NAVY = "1E3A5F"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(color="FFFFFF", bold=True)
SEV_FILL = {"critical": "F8D7DA", "high": "FDE2C8", "medium": "FFF3CD", "low": "E2F0D9"}


def _cell(v):
    if isinstance(v, (list, tuple, set)):
        return ", ".join(str(x) for x in v)[:32000]
    if isinstance(v, dict):
        return "; ".join(f"{k}: {x}" for k, x in v.items())[:32000]
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v  # neutralize spreadsheet formula injection from crawled text
    return v


def table(ws, rows: list[dict], cols: list[tuple[str, str]], start_row: int = 1, start_col: int = 1, widths=None):
    """Write a header + rows table; returns (first_data_row, last_row)."""
    for j, (_, label) in enumerate(cols):
        c = ws.cell(row=start_row, column=start_col + j, value=label)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, r in enumerate(rows, start=1):
        for j, (key, _) in enumerate(cols):
            c = ws.cell(row=start_row + i, column=start_col + j, value=_cell(r.get(key)))
            if key == "severity" and r.get(key) in SEV_FILL:
                c.fill = PatternFill("solid", fgColor=SEV_FILL[r[key]])
    for j, (key, label) in enumerate(cols):
        letter = get_column_letter(start_col + j)
        w = (widths or {}).get(key)
        if not w:
            sample = [len(str(_cell(r.get(key)) or "")) for r in rows[:200]] + [len(label)]
            w = min(max(sample) + 2, 70)
        ws.column_dimensions[letter].width = max(w, 8)
    if rows and start_row == 1:
        ws.auto_filter.ref = f"{get_column_letter(start_col)}1:{get_column_letter(start_col + len(cols) - 1)}{len(rows) + 1}"
        ws.freeze_panes = ws.cell(row=2, column=start_col)
    return start_row + 1, start_row + len(rows)


def _kv(ws, data: dict, start_row: int, col: int, title: tuple[str, str]):
    rows = [{"k": k, "v": v} for k, v in data.items()]
    return table(ws, rows, [("k", title[0]), ("v", title[1])], start_row=start_row, start_col=col)


def bar(ws, anchor, cats_ref, data_ref, title, horizontal=False, y_title=None, height=7.5, width=16):
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.title, ch.height, ch.width = title, height, width
    ch.add_data(data_ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    ch.legend = None
    if y_title:
        ch.y_axis.title = y_title
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    ws.add_chart(ch, anchor)


def pie(ws, anchor, cats_ref, data_ref, title):
    ch = PieChart()
    ch.title, ch.height, ch.width = title, 7.5, 11
    ch.add_data(data_ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showPercent = True
    ws.add_chart(ch, anchor)


def line(ws, anchor, cats_ref, data_refs, title, y_title=None, reverse=False):
    ch = LineChart()
    ch.title, ch.height, ch.width = title, 7.5, 18
    for ref in data_refs:
        ch.add_data(ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    if y_title:
        ch.y_axis.title = y_title
    if reverse:
        ch.y_axis.scaling.orientation = "maxMin"
    ws.add_chart(ch, anchor)


def _counter_block(ws, counts: dict, row: int, col: int, labels=("Item", "Count"), chart="bar", title="", anchor=None, horizontal=False):
    if not counts:
        return row
    first, last = _kv(ws, counts, row, col, labels)
    cats = Reference(ws, min_col=col, min_row=first, max_row=last)
    data = Reference(ws, min_col=col + 1, min_row=first - 1, max_row=last)
    anchor = anchor or f"{get_column_letter(col + 3)}{row}"
    if chart == "pie":
        pie(ws, anchor, cats, data, title)
    else:
        bar(ws, anchor, cats, data, title, horizontal=horizontal)
    return last + 2


def full_report_xlsx(result: dict, client: dict, site: dict, history: list[dict], search: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"SEO Audit: {site['url']}"
    ws["A1"].font = Font(size=16, bold=True, color=NAVY)
    ws["A2"] = f"Client: {client.get('name')} ({client.get('company') or ''}) · Generated {result['meta']['generated_at']}"
    ws["A3"] = f"Overall score: {result['scores']['overall']}/100 · Pages crawled: {result['meta']['pages_crawled']}"
    ws["A3"].font = Font(bold=True)
    nxt = _counter_block(ws, result["scores"]["categories"], 5, 1, ("Category", "Score"), title="Score by category", anchor="D5", horizontal=True)
    sev = Counter(i["severity"] for i in result["issues"])
    nxt = _counter_block(ws, {k: sev.get(k, 0) for k in ("critical", "high", "medium", "low")}, max(nxt, 22), 1, ("Severity", "Issues"), chart="pie", title="Issues by severity", anchor=f"D{max(nxt, 22)}")
    if history:
        r0 = max(nxt, 38)
        first, last = table(ws, history, [("started_at", "Audit date"), ("score", "Score"), ("pages_crawled", "Pages"),
                                          ("issues_critical", "Critical"), ("issues_high", "High")], start_row=r0, start_col=1)
        line(ws, f"H{r0}", Reference(ws, min_col=1, min_row=first, max_row=last),
             [Reference(ws, min_col=2, min_row=first - 1, max_row=last)], "Audit score over time", "Score")
    ws.column_dimensions["A"].width = 26

    ws = wb.create_sheet("Top 10 Actions")
    table(ws, result["top10"], [("rank", "#"), ("title", "Action"), ("severity", "Severity"), ("category", "Category"),
                                ("affected", "Affected"), ("why", "Why it matters"), ("how", "How to fix"), ("owner", "Owner"),
                                ("effort", "Effort"), ("impact", "Impact"), ("example_urls", "Example URLs")],
          widths={"why": 60, "how": 70, "example_urls": 60})
    first = 2
    last = 1 + len(result["top10"])
    if result["top10"]:
        bar(ws, "M2", Reference(ws, min_col=2, min_row=first, max_row=last),
            Reference(ws, min_col=5, min_row=1, max_row=last), "Affected URLs per priority action", horizontal=True, height=10)

    ws = wb.create_sheet("Roadmap 30-60-90")
    rows = []
    for ph in result["roadmap"]["phases"]:
        for t in ph["tasks"]:
            rows.append({"phase": ph["name"], **t})
    table(ws, rows, [("phase", "Phase"), ("task", "Task"), ("action", "Action"), ("owner", "Owner"), ("severity", "Severity"),
                     ("category", "Category"), ("affected", "Affected")], widths={"action": 80})
    _counter_block(ws, {ph["name"]: len(ph["tasks"]) for ph in result["roadmap"]["phases"]}, 1, 10, ("Phase", "Tasks"), title="Tasks per phase")

    ws = wb.create_sheet("All Issues")
    table(ws, result["issues"], [("severity", "Severity"), ("category", "Category"), ("title", "Issue"), ("count", "Affected"),
                                 ("description", "Description"), ("recommendation", "Recommendation"), ("urls", "URLs")],
          widths={"description": 60, "recommendation": 70, "urls": 80})
    cat = Counter()
    for i in result["issues"]:
        cat[i["category"]] += 1
    _counter_block(ws, dict(cat.most_common()), 1, 9, ("Category", "Issues"), title="Issues per category", horizontal=True)

    ws = wb.create_sheet("Pages")
    table(ws, result["pages"], [("url", "URL"), ("status", "Status"), ("indexable", "Indexable"), ("title", "Title"),
                                ("meta_description", "Meta description"), ("h1", "H1"), ("words", "Words"), ("inlinks", "Inlinks"),
                                ("outlinks", "Outlinks"), ("depth", "Depth"), ("ms", "Response ms"), ("size_kb", "HTML KB"),
                                ("canonical", "Canonical"), ("schema", "Schema"), ("primary_keyword", "Primary keyword"),
                                ("page_type", "Page type"), ("aio_score", "AIO score"), ("in_sitemap", "In sitemap")],
          widths={"title": 50, "meta_description": 60, "url": 60})
    _counter_block(ws, result["technical"]["status_counts"], 1, 21, ("Status", "Pages"), chart="pie", title="HTTP status codes")

    ws = wb.create_sheet("404s & 301 Plan")
    redir = result["technical"]["redirect_suggestions"]
    table(ws, redir, [("from", "Broken URL"), ("status", "Status"), ("to", "Suggested 301 target"), ("confidence", "Match confidence"),
                      ("linked_from", "Linked from")], widths={"from": 60, "to": 60, "linked_from": 80})

    ws = wb.create_sheet("Broken Links")
    table(ws, result["technical"]["broken_links"], [("source", "Source URL"), ("target", "Broken target"), ("status", "Status"),
                                                    ("anchor", "Anchor"), ("internal", "Internal")])

    ws = wb.create_sheet("Meta Suggestions")
    table(ws, result["meta_suggestions"], [("source_url", "Source URL"), ("target_keyword", "Target keyword"),
                                           ("current_title", "Current title"), ("suggested_title", "Suggested title"),
                                           ("current_description", "Current description"), ("suggested_description", "Suggested description"),
                                           ("secondary_keywords", "Secondary keywords"), ("reasons", "Why")],
          widths={"current_title": 45, "suggested_title": 45, "current_description": 60, "suggested_description": 60})

    _internal_link_sheets(wb, result)

    ws = wb.create_sheet("URL Structure")
    nxt = _counter_block(ws, result["links"]["stats"]["folders"], 1, 1, ("Folder", "URLs"), title="URLs per top-level folder", horizontal=True)
    table(ws, result["links"]["stats"]["url_problems"], [("url", "URL"), ("problems", "Problems")], start_row=max(nxt, 20), start_col=1)

    ws = wb.create_sheet("Schema")
    nxt = _counter_block(ws, result["technical"]["schema_types"] or {"(none)": 0}, 1, 1, ("Schema type", "Count"), title="Schema types found", horizontal=True)
    r0 = max(nxt, 18)
    first, last = table(ws, result["technical"]["schema_suggestions"], [("url", "Page without structured data"), ("suggested", "Suggested schema")], start_row=r0, start_col=1)
    table(ws, result["technical"]["schema_problems"], [("url", "URL"), ("type", "Type"), ("problems", "Validation problems")], start_row=last + 3, start_col=1)

    kw = result["keywords"]
    ws = wb.create_sheet("Keyword Research")
    table(ws, kw["keywords"], [("keyword", "Keyword"), ("intent", "Intent"), ("funnel", "Funnel"), ("priority", "Priority"),
                               ("status", "Status"), ("target_url", "Target URL"), ("suggested_slug", "Suggested slug (gaps)"),
                               ("search_volume", "Volume"), ("cpc", "CPC"), ("gsc_impressions", "GSC impr."), ("gsc_clicks", "GSC clicks"),
                               ("gsc_position", "GSC pos."), ("source", "Source"), ("seed", "Cluster")])
    nxt = _counter_block(ws, kw["intent_counts"], 1, 16, ("Intent", "Keywords"), chart="pie", title="Keywords by search intent")
    _counter_block(ws, kw["status_counts"], nxt + 12, 16, ("Status", "Keywords"), title="Coverage: covered vs content gaps")

    ws = wb.create_sheet("Keyword Clusters")
    table(ws, kw["clusters"], [("cluster", "Cluster"), ("size", "Keywords"), ("pillar_url", "Pillar URL"), ("keywords", "Members")], widths={"keywords": 100})
    if kw["clusters"]:
        n = min(len(kw["clusters"]), 15)
        bar(ws, "F2", Reference(ws, min_col=1, min_row=2, max_row=1 + n), Reference(ws, min_col=2, min_row=1, max_row=1 + n), "Cluster size", horizontal=True, height=10)

    ws = wb.create_sheet("Search Performance")
    if search.get("has_data"):
        first, last = table(ws, search["daily"], [("date", "Date"), ("clicks", "Clicks"), ("impressions", "Impressions"), ("position", "Avg position")])
        cats = Reference(ws, min_col=1, min_row=first, max_row=last)
        line(ws, "F2", cats, [Reference(ws, min_col=2, min_row=1, max_row=last), Reference(ws, min_col=3, min_row=1, max_row=last)], "Organic clicks & impressions")
        line(ws, "F18", cats, [Reference(ws, min_col=4, min_row=1, max_row=last)], "Average position (lower is better)", reverse=True)
        r0 = 36
        ws.cell(row=r0 - 1, column=6, value="Ranking improvements (current vs previous 28 days)").font = Font(bold=True)
        f2, l2 = table(ws, search["improved"][:25], [("query", "Query"), ("prev_position", "Prev pos"), ("position", "Pos"), ("change", "Gain"), ("clicks", "Clicks")], start_row=r0, start_col=6)
        if search["improved"]:
            bar(ws, f"M{r0}", Reference(ws, min_col=6, min_row=f2, max_row=min(l2, f2 + 14)), Reference(ws, min_col=9, min_row=f2 - 1, max_row=min(l2, f2 + 14)), "Top ranking gains (positions)", horizontal=True, height=10)
        r1 = l2 + 3
        ws.cell(row=r1 - 1, column=6, value="Ranking declines").font = Font(bold=True)
        table(ws, search["declined"][:25], [("query", "Query"), ("prev_position", "Prev pos"), ("position", "Pos"), ("change", "Change"), ("clicks", "Clicks")], start_row=r1, start_col=6)
        _counter_block(ws, search["position_buckets"], r1 + 30, 6, ("Position range", "Keywords"), title="Keywords by ranking position")
    else:
        ws["A1"] = "No Search Console data yet. Add a Windsor.ai API key in Settings and click 'Sync GSC/GA4'."

    ws = wb.create_sheet("Core Web Vitals")
    rows = []
    for r in result["performance"]["results"]:
        if "error" in r:
            rows.append({"url": r["url"], "strategy": r["strategy"], "error": r["error"]})
            continue
        rows.append({"url": r["url"], "strategy": r["strategy"], "performance": r["scores"].get("performance"), "seo": r["scores"].get("seo"),
                     "lcp": r["field"].get("LCP") or r["lab"].get("LCP"), "inp": r["field"].get("INP"), "cls": r["field"].get("CLS") if r["field"].get("CLS") is not None else r["lab"].get("CLS"),
                     "ttfb": r["lab"].get("TTFB"), "weight": r["page_weight_kb"], "top": [o["title"] for o in r["opportunities"][:3]]})
    first, last = table(ws, rows, [("url", "URL"), ("strategy", "Device"), ("performance", "Perf score"), ("seo", "SEO score"), ("lcp", "LCP ms"),
                                   ("inp", "INP ms"), ("cls", "CLS"), ("ttfb", "TTFB ms"), ("weight", "Weight KB"), ("top", "Top opportunities"), ("error", "Error")])
    if rows and any("performance" in r for r in rows):
        bar(ws, "M2", Reference(ws, min_col=1, min_row=first, max_row=last), Reference(ws, min_col=3, min_row=1, max_row=last), "Lighthouse performance score", horizontal=True)

    ws = wb.create_sheet("AIO & GEO")
    g = result["geo"]
    ws["A1"] = g["brief"]["summary"]
    ws["A1"].font = Font(bold=True)
    first, last = table(ws, g["pages"], [("url", "URL"), ("score", "AIO score"), ("strengths", "Strengths"), ("fixes", "Fixes")], start_row=3, start_col=1,
                        widths={"strengths": 50, "fixes": 90})
    buckets = Counter("0-24" if p["score"] < 25 else "25-49" if p["score"] < 50 else "50-74" if p["score"] < 75 else "75-100" for p in g["pages"])
    _counter_block(ws, {k: buckets.get(k, 0) for k in ("0-24", "25-49", "50-74", "75-100")}, 3, 7, ("Score band", "Pages"), title="AI answer-readiness distribution")
    table(ws, g["brief"]["prerequisites"], [("check", "Prerequisite"), ("ok", "Met"), ("detail", "Detail")], start_row=last + 3, start_col=1)

    ws = wb.create_sheet("Conversion Tracking")
    conv = result["conversion"]
    first, last = table(ws, conv["events"], [("priority", "Priority"), ("page_url", "Page URL"), ("event_name", "GA4 event"), ("element", "Element"),
                                             ("gtm_trigger", "GTM trigger"), ("parameters", "Parameters"), ("key_event", "Key event"), ("why", "Why")],
                        widths={"gtm_trigger": 55, "element": 45})
    ev = Counter(e["event_name"] for e in conv["events"])
    nxt = _counter_block(ws, dict(ev.most_common(15)), 1, 10, ("Event", "Placements"), title="Recommended events", horizontal=True)
    _counter_block(ws, conv["page_types"], nxt + 8, 10, ("Page type", "Pages"), chart="pie", title="Pages by commercial role")
    ws2 = wb.create_sheet("Lead-Gen Opportunities")
    f1, l1 = table(ws2, conv["lead_gaps"], [("url", "URL"), ("page_type", "Page type"), ("problem", "Problem"), ("recommendation", "Recommendation")])
    table(ws2, conv["opportunities"], [("url", "URL"), ("page_type", "Page type"), ("opportunity", "Opportunity"), ("recommendation", "Recommendation"), ("sessions", "Sessions")], start_row=l1 + 3)

    ws = wb.create_sheet("Off-page & Guest Posts")
    off = result["offpage"]
    f1, l1 = table(ws, off["platforms"], [("name", "Platform"), ("url", "URL"), ("cost", "Cost"), ("type", "Type"), ("notes", "Notes")])
    f2, l2 = table(ws, off["pr_sources"], [("name", "Digital PR source"), ("url", "URL"), ("cost", "Cost"), ("type", "Type")], start_row=l1 + 3)
    table(ws, off["footprints"], [("keyword", "Keyword"), ("type", "Type"), ("query", "Search footprint"), ("google_url", "Google search")], start_row=l2 + 3)
    _counter_block(ws, {"Free": off["free_count"], "Paid / membership": off["paid_count"]}, 1, 8, ("Cost", "Platforms"), chart="pie", title="Free vs paid platforms")

    ws = wb.create_sheet("Robots, Sitemaps, llms.txt")
    rb = result["robots"]
    _kv(ws, {"robots.txt": rb.get("url"), "exists": rb.get("exists"), "status": rb.get("status"), "disallow all": rb.get("disallow_all"),
             "sitemaps declared": rb.get("sitemaps"), "sitemap URLs found": result["sitemaps"].get("url_count"),
             "llms.txt": result["llms_txt"].get("llms.txt", {}).get("exists"), "llms-full.txt": result["llms_txt"].get("llms-full.txt", {}).get("exists")},
        1, 1, ("Check", "Value"))
    _kv(ws, rb.get("ai_crawlers") or {}, 12, 1, ("AI crawler", "robots.txt"))
    table(ws, result["sitemaps"].get("checked", []), [("url", "Sitemap"), ("status", "Status"), ("type", "Type"), ("count", "Entries"), ("error", "Error")], start_row=26, start_col=1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _internal_link_sheets(wb, result: dict):
    lk = result["links"]
    ws = wb.create_sheet("Internal Link Suggestions")
    first, last = table(ws, lk["suggestions"], [("priority", "Priority"), ("source_url", "Source URL (add link here)"), ("target_url", "Target URL (link to)"),
                                                ("anchor_text", "Suggested anchor"), ("context", "Context in source"), ("target_inlinks", "Target inlinks now"),
                                                ("reason", "Reason"), ("source_title", "Source title"), ("target_title", "Target title")],
                        widths={"context": 80, "source_url": 55, "target_url": 55})
    _counter_block(ws, dict(Counter(s["priority"] for s in lk["suggestions"])), 1, 11, ("Priority", "Suggestions"), chart="pie", title="Link suggestions by priority")
    ws = wb.create_sheet("Orphans & Inlinks")
    pages = sorted(result["pages"], key=lambda p: p.get("inlinks") or 0)
    first, last = table(ws, [p for p in pages if p.get("indexable")], [("url", "URL"), ("inlinks", "Internal inlinks"), ("outlinks", "Outlinks"),
                                                                        ("depth", "Click depth"), ("in_sitemap", "In sitemap"), ("primary_keyword", "Primary keyword")])
    _counter_block(ws, lk["stats"]["depth_distribution"], 1, 8, ("Click depth", "Pages"), title="Click depth distribution")
    top = lk["stats"]["top_linked"][:15]
    if top:
        f, l = table(ws, top, [("url", "Most linked URL"), ("inlinks", "Inlinks")], start_row=22, start_col=8)
        bar(ws, "K22", Reference(ws, min_col=8, min_row=f, max_row=l), Reference(ws, min_col=9, min_row=f - 1, max_row=l), "Most internally linked pages", horizontal=True, height=10)
    ws = wb.create_sheet("Links to Redirects")
    table(ws, result["technical"]["links_to_redirects"], [("source", "Source URL"), ("target", "Linked URL (redirects)"), ("final", "Update link to"), ("anchor", "Anchor")])


def internal_links_xlsx(result: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    _internal_link_sheets(wb, result)
    ws = wb.create_sheet("Broken Internal Links")
    table(ws, [b for b in result["technical"]["broken_links"] if b.get("internal")],
          [("source", "Source URL"), ("target", "Broken target"), ("status", "Status"), ("anchor", "Anchor")])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
